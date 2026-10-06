#!/usr/bin/env python3
"""excel_to_md.py — Converte Excel de volta para .md (uma tabela por aba)."""
import os, sys
from openpyxl import load_workbook

# Console do Windows costuma vir em cp1252 — garante que ✓ não quebre a execução
try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass

def sheet_to_md(ws, sheet_name):
    lines = [f"# {sheet_name}\n"]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return '\n'.join(lines)
    max_cols = max(len(r) for r in rows)
    table_rows = []
    for row in rows:
        cells = [str(cell) if cell is not None else 'nan' for cell in row]
        cells += ['nan'] * (max_cols - len(cells))
        table_rows.append('| ' + ' | '.join(cells) + ' |')
    lines.append(table_rows[0])
    lines.append('|' + '|'.join([':---'] * max_cols) + '|')
    lines.extend(table_rows[1:])
    return '\n'.join(lines) + '\n'

def excel_to_mds(excel_path, output_dir):
    wb = load_workbook(excel_path, read_only=True)
    os.makedirs(output_dir, exist_ok=True)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        md_content = sheet_to_md(ws, sheet_name)
        with open(os.path.join(output_dir, f"{sheet_name}.md"), 'w', encoding='utf-8') as f:
            f.write(md_content)
        print(f"  ✓ {sheet_name}.md")

if __name__ == '__main__':
    excel_to_mds(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else '.')
